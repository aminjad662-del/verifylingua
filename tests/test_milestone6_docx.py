import os
import io
import json
import zipfile
import pytest
import docx
from lxml import etree
from fastapi.testclient import TestClient

from api.main import app
from api.docx_pipeline import DocxPipeline, W_NS, NSMAP
from api.analyze import DocumentAnalysis
from api.render import reconstruct_document

CORPUS_DIR = os.path.join(os.path.dirname(__file__), "corpus")
FILE_11_PATH = os.path.join(CORPUS_DIR, "11_complex_elements.docx")
REPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "reports", "milestone-6")
VISUAL_DIR = os.path.join(REPORTS_DIR, "visual")

os.makedirs(VISUAL_DIR, exist_ok=True)

@pytest.fixture(scope="module")
def file11_bytes() -> bytes:
    with open(FILE_11_PATH, "rb") as f:
        return f.read()

@pytest.fixture(scope="module")
def api_client() -> TestClient:
    return TestClient(app)


def test_gate_6_1_extraction_and_analysis(file11_bytes):
    """Gate 6.1: Direct OpenXML extraction of paragraphs, tables, headers, and footers."""
    analysis = DocxPipeline.extract_docx_analysis(file11_bytes, "11_complex_elements.docx")
    
    assert isinstance(analysis, DocumentAnalysis)
    assert len(analysis.pages) == 1
    page = analysis.pages[0]
    
    # 1. Blocks extracted
    assert len(page.blocks) >= 10
    
    # 2. Check roles: header, heading, paragraph, table_cell, footer
    roles = {b.role for b in page.blocks}
    assert "header" in roles
    assert "heading" in roles
    assert "paragraph" in roles
    assert "table_cell" in roles
    assert "footer" in roles
    
    # 3. Check inline tag extraction on student paragraph
    student_blocks = [b for b in page.blocks if "Alejandro Garcia" in b.text]
    assert len(student_blocks) == 1
    sb = student_blocks[0]
    assert "<b1>Alejandro Garcia</b1>" in sb.text
    assert "<i2>Highest Academic Honors</i2>" in sb.text
    
    # 4. Check table structure
    assert len(page.tables) == 1
    tbl = page.tables[0]
    assert tbl.rows_count == 3
    assert tbl.cols_count == 3
    assert len(tbl.cells) == 9
    assert tbl.cells[0].text == "Course Code"
    assert tbl.cells[0].is_header is True


def test_gate_6_2_reconstruction_spanish_formatting(file11_bytes):
    """Gate 6.2: Reconstruction preserves run-level bold/italic and table structure in Spanish."""
    analysis = DocxPipeline.extract_docx_analysis(file11_bytes, "11_complex_elements.docx")
    page = analysis.pages[0]
    
    # Build translated segments
    spanish_translations = {
        "OFFICIAL ACADEMIC TRANSCRIPT": "CERTIFICADO ACADÉMICO OFICIAL",
        "Page 1 of 1 - Confidential": "Página 1 de 1 - Confidencial",
        "Student Name: <b1>Alejandro Garcia</b1> has completed all requirements with <i2>Highest Academic Honors</i2>.":
            "Nombre del estudiante: <b1>Alejandro Garcia</b1> ha cumplido todos los requisitos con <i2>los más altos honores académicos</i2>.",
        "This document certifies that the individual named above has demonstrated exemplary performance in the following curriculum:":
            "El presente documento certifica que la persona antes mencionada ha demostrado un rendimiento ejemplar en el siguiente plan de estudios:",
        "Course Code": "Código",
        "Course Title": "Título del Curso",
        "Final Grade": "Calificación",
        "MATH-401": "MATH-401",
        "Advanced Mathematics": "Matemáticas Avanzadas",
        "A (Honors)": "A (Con Honores)",
        "PHYS-302": "PHYS-302",
        "Theoretical Physics": "Física Teórica",
        "A": "A"
    }

    segments = []
    for b in page.blocks:
        trans = spanish_translations.get(b.text, f"[ES] {b.text}")
        segments.append({
            "id": b.id,
            "block_id": b.id,
            "page_number": 1,
            "source_text": b.text,
            "translated_text": trans
        })

    rendered_bytes, preview_bytes, qa_records = reconstruct_document(
        original_bytes=file11_bytes,
        file_format="docx",
        analysis=analysis,
        translated_segments=segments,
        target_lang="es"
    )

    assert len(rendered_bytes) > 1000
    assert len(preview_bytes) > 1000
    assert len(qa_records) == 1
    
    # Save output for inspection
    out_docx_path = os.path.join(REPORTS_DIR, "11_complex_elements_translated_es.docx")
    with open(out_docx_path, "wb") as f:
        f.write(rendered_bytes)
    assert os.path.exists(out_docx_path)

    # Save preview image
    preview_img_path = os.path.join(VISUAL_DIR, "11_complex_elements_es_preview.jpg")
    with open(preview_img_path, "wb") as f:
        f.write(preview_bytes)
    assert os.path.exists(preview_img_path)

    # Verify document structure with python-docx
    doc = docx.Document(io.BytesIO(rendered_bytes))
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 3
    assert table.rows[0].cells[0].text == "Código"
    assert table.rows[1].cells[1].text == "Matemáticas Avanzadas"

    # Verify bold and italic runs in paragraph 1
    p1 = doc.paragraphs[1]
    assert "Alejandro Garcia" in p1.text
    bold_runs = [r for r in p1.runs if r.bold is True]
    assert len(bold_runs) == 1
    assert bold_runs[0].text == "Alejandro Garcia"

    italic_runs = [r for r in p1.runs if r.italic is True]
    assert len(italic_runs) == 1
    assert "los más altos honores académicos" in italic_runs[0].text

    # Validation check
    val_report = qa_records[0]["validation"]
    assert val_report["is_valid_zip"] is True
    assert val_report["xml_schema_valid"] is True
    assert val_report["tables_preserved"] is True


def test_gate_6_3_arabic_bidi_and_table_visual(file11_bytes):
    """Gate 6.3: Arabic BiDi shaping, w:bidi, w:rtl, w:bidiVisual, and complex-script fonts."""
    analysis = DocxPipeline.extract_docx_analysis(file11_bytes, "11_complex_elements.docx")
    page = analysis.pages[0]

    ar_translations = {
        "OFFICIAL ACADEMIC TRANSCRIPT": "السجل الأكاديمي الرسمي",
        "Page 1 of 1 - Confidential": "صفحة 1 من 1 - سري للغاية",
        "Student Name: <b1>Alejandro Garcia</b1> has completed all requirements with <i2>Highest Academic Honors</i2>.":
            "اسم الطالب: <b1>Alejandro Garcia</b1> قد أتم جميع المتطلبات مع <i2>أعلى مرتبة شرف أكاديمية</i2>.",
        "This document certifies that the individual named above has demonstrated exemplary performance in the following curriculum:":
            "تشهد هذه الوثيقة بأن الفرد المذكور أعلاه قد أظهر أداءً نموذجياً في المنهج التالي:",
        "Course Code": "رمز المقرر",
        "Course Title": "اسم المقرر",
        "Final Grade": "الدرجة النهائية",
        "MATH-401": "MATH-401",
        "Advanced Mathematics": "الرياضيات المتقدمة",
        "A (Honors)": "A (مرتبة الشرف)",
        "PHYS-302": "PHYS-302",
        "Theoretical Physics": "الفيزياء النظرية",
        "A": "A"
    }

    segments = [
        {
            "id": b.id,
            "block_id": b.id,
            "page_number": 1,
            "source_text": b.text,
            "translated_text": ar_translations.get(b.text, f"ترجمة: {b.text}")
        }
        for b in page.blocks
    ]

    rendered_bytes, preview_bytes, qa_records = reconstruct_document(
        original_bytes=file11_bytes,
        file_format="docx",
        analysis=analysis,
        translated_segments=segments,
        target_lang="ar"
    )

    out_ar_path = os.path.join(REPORTS_DIR, "11_complex_elements_translated_ar.docx")
    with open(out_ar_path, "wb") as f:
        f.write(rendered_bytes)
    assert os.path.exists(out_ar_path)

    ar_preview_path = os.path.join(VISUAL_DIR, "11_complex_elements_ar_preview.jpg")
    with open(ar_preview_path, "wb") as f:
        f.write(preview_bytes)
    assert os.path.exists(ar_preview_path)

    # Inspect XML directly with lxml
    zf = zipfile.ZipFile(io.BytesIO(rendered_bytes), "r")
    tree = etree.fromstring(zf.read("word/document.xml"))

    # Verify w:bidi on paragraphs
    bidi_paragraphs = tree.xpath(".//w:pPr/w:bidi", namespaces=NSMAP)
    assert len(bidi_paragraphs) >= 1

    # Verify w:rtl on runs
    rtl_runs = tree.xpath(".//w:rPr/w:rtl", namespaces=NSMAP)
    assert len(rtl_runs) >= 1

    # Verify w:rFonts complex script
    cs_fonts = tree.xpath(".//w:rPr/w:rFonts[@w:cs='Arial']", namespaces=NSMAP)
    assert len(cs_fonts) >= 1

    # Verify w:bidiVisual on tables
    bidi_tables = tree.xpath(".//w:tblPr/w:bidiVisual", namespaces=NSMAP)
    assert len(bidi_tables) == 1

    # Verify python-docx loads without error
    doc = docx.Document(io.BytesIO(rendered_bytes))
    assert len(doc.tables) == 1
    assert len(doc.paragraphs) >= 3


def test_gate_6_4_full_api_workflow(api_client, file11_bytes):
    """Gate 6.4: Full API pipeline intake -> analyze -> translate -> render -> preview -> download."""
    # 1. Create Job
    res = api_client.post("/api/jobs", json={
        "sourceFilename": "11_complex_elements.docx",
        "sourceFormat": "docx",
        "sourceLanguage": "en",
        "targetLanguage": "es"
    })
    assert res.status_code == 200
    job_id = res.json()["id"]

    # 2. Intake Upload
    res = api_client.post(
        f"/api/jobs/{job_id}/intake",
        files={"file": ("11_complex_elements.docx", file11_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert res.status_code == 200
    assert res.json()["sourceFormat"] == "docx"

    # 3. Analyze
    res = api_client.post(f"/api/jobs/{job_id}/analyze")
    assert res.status_code == 200
    assert res.json()["currentStage"] == "translating"

    # 4. Translate
    res = api_client.post(f"/api/jobs/{job_id}/translate", json={})
    assert res.status_code == 200
    assert res.json()["currentStage"] == "rendering"

    # 5. Render
    res = api_client.post(f"/api/jobs/{job_id}/render")
    assert res.status_code == 200
    assert res.json()["currentStage"] == "qa"

    # 6. Preview
    res = api_client.get(f"/api/jobs/{job_id}/preview")
    assert res.status_code == 200
    assert res.headers["content-type"] == "image/jpeg"
    assert len(res.content) > 1000

    # 7. Download
    res = api_client.get(f"/api/jobs/{job_id}/download")
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in res.headers["content-type"]
    assert len(res.content) > 1000

    # Verify downloaded file opens cleanly
    doc = docx.Document(io.BytesIO(res.content))
    assert len(doc.tables) == 1
    assert len(doc.paragraphs) >= 3


def test_gate_6_5_non_standard_minimal_docx():
    """Gate 6.5: Universal OpenXML support for minimal DOCX zips without _rels."""
    minimal_buf = io.BytesIO()
    with zipfile.ZipFile(minimal_buf, "w") as zf:
        zf.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="application/xml"/></Types>')
        zf.writestr("word/header1.xml", '<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:t>Header Text</w:t></w:r></w:p></w:hdr>')
        zf.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Body Text</w:t></w:r></w:p></w:body></w:document>')
    
    raw_bytes = minimal_buf.getvalue()
    analysis = DocxPipeline.extract_docx_analysis(raw_bytes, "minimal.docx")
    assert len(analysis.pages[0].blocks) == 2
    
    recon_bytes = DocxPipeline.reconstruct_docx(raw_bytes, [
        {"block_id": "header1_p_0", "translated_text": "Texto del Encabezado"},
        {"block_id": "document_p_0", "translated_text": "Texto del Cuerpo"}
    ], target_lang="es")
    
    assert len(recon_bytes) > 0
    recon_zf = zipfile.ZipFile(io.BytesIO(recon_bytes), "r")
    assert "word/document.xml" in recon_zf.namelist()
    doc_xml = recon_zf.read("word/document.xml").decode("utf-8")
    assert "Texto del Cuerpo" in doc_xml
